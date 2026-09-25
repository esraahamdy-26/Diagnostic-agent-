"""
Multimodal/document ingestion for the conversational diagnostic agent.

The frontend can attach an image/PDF/document and type a question in the same
message. This module turns supported attachments into useful context for the
conversation:
  - images: Gemini vision reads visible text + useful document observations
  - PDFs: extract embedded text first; for scanned PDFs, render pages and use vision
  - txt/csv: read text directly
  - docx: extract paragraphs and table cells

Important: extraction is not a diagnosis. The conversational layer is still
responsible for patient-facing explanation and safety routing.
"""

from __future__ import annotations

import base64
import io
import os
from pathlib import Path
from typing import Any


class VisionExtractionError(Exception):
    pass


def _provider_and_key(api_key: str | None, provider: str | None) -> tuple[str | None, str | None]:
    if provider:
        return provider, api_key or os.getenv("GEMINI_API_KEY") or os.getenv("OPENAI_API_KEY")
    if api_key:
        return ("openai" if api_key.startswith("sk-") else "gemini"), api_key
    if os.getenv("GEMINI_API_KEY"):
        return "gemini", os.getenv("GEMINI_API_KEY")
    if os.getenv("OPENAI_API_KEY"):
        return "openai", os.getenv("OPENAI_API_KEY")
    return None, None


def _vision_prompt() -> str:
    return (
        "You are reading an attachment supplied by a patient.\n"
        "First identify what kind of document/image it appears to be.\n"
        "Transcribe visible text and numeric values accurately, preserving units and reference ranges.\n"
        "If it is a medical report or lab result, organize the visible tests/findings clearly.\n"
        "If there are parts you cannot read, say [غير واضح] instead of guessing.\n"
        "Do not diagnose the patient and do not invent missing values.\n"
        "Return plain text only. Use the same language as the document when possible; English medical "
        "terms are okay when they are the actual names on the document."
    )


def _gemini_image(image_bytes: bytes, mime_type: str, api_key: str) -> str:
    import google.generativeai as genai
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(os.getenv("GEMINI_MODEL", "gemini-3.6-flash"))
    response = model.generate_content([
        _vision_prompt(),
        {"mime_type": mime_type, "data": image_bytes},
    ])
    text = (getattr(response, "text", "") or "").strip()
    if not text:
        raise VisionExtractionError("الموديل قدر يستقبل الصورة لكن مارجعش نص مقروء منها.")
    return text


def _openai_image(image_bytes: bytes, mime_type: str, api_key: str) -> str:
    from openai import OpenAI
    client = OpenAI(api_key=api_key, timeout=30.0)
    b64 = base64.b64encode(image_bytes).decode("utf-8")
    response = client.chat.completions.create(
        model=os.getenv("OPENAI_VISION_MODEL", os.getenv("OPENAI_MODEL", "gpt-4o-mini")),
        messages=[{
            "role": "user",
            "content": [
                {"type": "text", "text": _vision_prompt()},
                {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64}"}},
            ],
        }],
    )
    text = (response.choices[0].message.content or "").strip()
    if not text:
        raise VisionExtractionError("الموديل قدر يستقبل الصورة لكن مارجعش نص مقروء منها.")
    return text


def _extract_pdf(pdf_bytes: bytes, api_key: str | None, provider: str | None) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise VisionExtractionError("لازم تثبتي PyMuPDF عشان الـ PDF يتقري.") from exc

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages_text: list[str] = []
    scanned_pages: list[bytes] = []

    for page in doc:
        text = (page.get_text("text") or "").strip()
        if text:
            pages_text.append(text)
        else:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            scanned_pages.append(pix.tobytes("png"))

    # If the PDF already contains selectable text, preserve it exactly enough
    # for downstream RAG/NLU. For scanned pages, use vision page-by-page.
    if not scanned_pages:
        return "\n\n--- PDF page ---\n\n".join(pages_text) or "[PDF فارغ أو غير قابل للقراءة]"

    provider, api_key = _provider_and_key(api_key, provider)
    if not api_key or not provider:
        if pages_text:
            return "\n\n--- PDF page ---\n\n".join(pages_text)
        raise VisionExtractionError("الـ PDF ده سكان/صورة ومحتاج Vision API عشان أقدر أقرأه.")

    extracted = list(pages_text)
    for idx, image_bytes in enumerate(scanned_pages, start=1):
        if provider == "gemini":
            page_text = _gemini_image(image_bytes, "image/png", api_key)
        elif provider == "openai":
            page_text = _openai_image(image_bytes, "image/png", api_key)
        else:
            raise VisionExtractionError(f"Unknown provider: {provider}")
        extracted.append(f"[Scanned PDF page {idx}]\n{page_text}")
    return "\n\n--- PDF page ---\n\n".join(extracted)


def _extract_plain_document(content: bytes, mime_type: str, filename: str) -> str | None:
    ext = Path(filename or "").suffix.lower()
    if ext in {".txt", ".md", ".csv", ".json"} or mime_type.startswith("text/"):
        return content.decode("utf-8", errors="replace").strip()

    if ext == ".docx" or "wordprocessingml" in mime_type:
        try:
            from docx import Document
        except ImportError as exc:
            raise VisionExtractionError("لازم تثبتي python-docx عشان ملفات Word تتقري.") from exc
        doc = Document(io.BytesIO(content))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip() or "[ملف Word فاضي أو مفيهوش نص قابل للقراءة]"
    return None


def extract_attachment_content(
    content: bytes,
    mime_type: str = "application/octet-stream",
    filename: str = "attachment",
    api_key: str | None = None,
    provider: str | None = None,
) -> dict[str, Any]:
    """Return normalized attachment context for the conversational agent."""
    if not content:
        raise VisionExtractionError("الملف فاضي.")

    mime = (mime_type or "").lower()
    ext = Path(filename or "").suffix.lower()

    direct = _extract_plain_document(content, mime, filename)
    if direct is not None:
        return {"kind": "document", "filename": filename, "mime_type": mime, "text": direct}

    if mime == "application/pdf" or ext == ".pdf":
        text = _extract_pdf(content, api_key, provider)
        return {"kind": "pdf", "filename": filename, "mime_type": mime, "text": text}

    if mime.startswith("image/") or ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}:
        provider, api_key = _provider_and_key(api_key, provider)
        if not provider or not api_key:
            raise VisionExtractionError("محتاجين GEMINI_API_KEY أو OPENAI_API_KEY عشان أقرأ الصور.")
        if provider == "gemini":
            text = _gemini_image(content, mime or "image/jpeg", api_key)
        elif provider == "openai":
            text = _openai_image(content, mime or "image/jpeg", api_key)
        else:
            raise VisionExtractionError(f"Unknown provider: {provider}")
        return {"kind": "image", "filename": filename, "mime_type": mime, "text": text}

    raise VisionExtractionError(
        "نوع الملف ده مش مدعوم حالياً. ارفعي صورة، PDF، TXT/CSV/JSON أو Word (.docx)."
    )


def extract_text_from_image(
    image_bytes: bytes,
    mime_type: str = "image/jpeg",
    api_key: str | None = None,
    provider: str | None = None,
) -> str:
    """Backward-compatible wrapper used by older callers."""
    return extract_attachment_content(
        image_bytes,
        mime_type=mime_type,
        filename="image",
        api_key=api_key,
        provider=provider,
    )["text"]
